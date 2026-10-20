"""강의 슬라이드 공통 도우미 (python-pptx)

모든 강의 빌드 스크립트에서 import 해서 사용.
색·글꼴·레이아웃을 한곳에서 관리해 강의 간 디자인을 통일.
"""
import json
import os

import av
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FONT = "Malgun Gothic"
BG = RGBColor(0xFA, 0xFA, 0xF7)
INK = RGBColor(0x1F, 0x2A, 0x44)
BLUE = RGBColor(0x2E, 0x86, 0xAB)
ORANGE = RGBColor(0xE0, 0x7A, 0x2E)
GREEN = RGBColor(0x3A, 0x9D, 0x5D)
GRAY = RGBColor(0x8A, 0x8F, 0x98)
LIGHT = RGBColor(0xEE, 0xEC, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = Inches(13.333), Inches(7.5)


def _font(run, size, color=INK, bold=False):
    f = run.font
    f.name, f.size, f.bold = FONT, Pt(size), bold
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text(slide, x, y, w, h, lines, size=22, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullet=None, spacing=1.15, para_space=8):
    """lines: 문자열 또는 (문자열, {size,color,bold}) 목록. '**' 로 감싼 부분은 강조색 굵게."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        opt = {}
        if isinstance(line, tuple):
            line, opt = line
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opt.get("align", align)
        p.line_spacing = spacing
        p.space_after = Pt(para_space)
        prefix = opt.get("bullet", bullet)
        parts = line.split("**")
        if prefix:
            parts[0] = prefix + parts[0]
        for j, part in enumerate(parts):
            if not part:
                continue
            r = p.add_run()
            r.text = part
            hi = j % 2 == 1
            _font(r, opt.get("size", size),
                  opt.get("hicolor", ORANGE) if hi else opt.get("color", color),
                  True if hi else opt.get("bold", bold))
    return tb


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.5)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def picture_cover(slide, path, x, y, w, h):
    """상자를 꽉 채우도록 가운데 기준으로 잘라서 넣기"""
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(path, x, y, w, h)
    box, img = w / h, iw / ih
    if img > box:
        c = (1 - box / img) / 2
        pic.crop_left = pic.crop_right = c
    else:
        c = (1 - img / box) / 2
        pic.crop_top = pic.crop_bottom = c
    return pic


def poster(video, out_png):
    """영상의 마지막 프레임을 포스터 이미지로 저장"""
    if not os.path.exists(out_png):
        last = None
        for fr in av.open(video).decode(video=0):
            last = fr
        os.makedirs(os.path.dirname(out_png), exist_ok=True)
        last.to_image().save(out_png)
    return out_png


class Deck:
    def __init__(self, footer, image_dir=None):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SW, SH
        self.footer = footer
        self.n = 0
        self.credits = {}
        if image_dir and os.path.exists(os.path.join(image_dir, "credits.json")):
            self.credits = json.load(open(os.path.join(image_dir, "credits.json"), encoding="utf-8"))
        self.image_dir = image_dir

    # ---------- 기본 틀 ----------
    def _new(self, notes=None, dark=False):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = INK if dark else BG
        self.n += 1
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def _chrome(self, s, title):
        rect(s, Inches(0.5), Inches(0.42), Inches(0.12), Inches(0.62), ORANGE)
        text(s, Inches(0.75), Inches(0.3), Inches(11.8), Inches(0.85), title, 30, INK, True,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(0.5), Inches(7.02), Inches(8), Inches(0.35), self.footer, 11, GRAY)
        text(s, Inches(12.0), Inches(7.02), Inches(0.83), Inches(0.35), str(self.n), 11, GRAY,
             align=PP_ALIGN.RIGHT)

    def _credit(self, s, img_names, x=Inches(4.5), w=Inches(7.4)):
        parts = []
        for name in img_names:
            c = self.credits.get(name)
            if c and "Public domain" not in c["license"]:
                parts.append(f"사진: {c['artist']} / {c['license']} / Wikimedia Commons")
        if parts:
            text(s, x, Inches(7.02), w, Inches(0.35), " · ".join(parts), 9, GRAY, align=PP_ALIGN.RIGHT)

    def img(self, name):
        return os.path.join(self.image_dir, name)

    # ---------- 슬라이드 유형 ----------
    def title(self, title, subtitle, small, notes=None):
        s = self._new(notes, dark=True)
        rect(s, Inches(0.9), Inches(2.55), Inches(0.16), Inches(1.6), ORANGE)
        text(s, Inches(1.3), Inches(2.3), Inches(11), Inches(1.3), title, 54, WHITE, True)
        text(s, Inches(1.3), Inches(3.55), Inches(11), Inches(0.8), subtitle, 26, RGBColor(0xC9, 0xD3, 0xE3))
        text(s, Inches(1.3), Inches(5.9), Inches(11), Inches(0.6), small, 16, GRAY)
        return s

    def section(self, num, title, desc, notes=None):
        s = self._new(notes, dark=True)
        text(s, Inches(1.0), Inches(2.0), Inches(3), Inches(1.6), f"{num:02d}", 96, ORANGE, True)
        text(s, Inches(1.0), Inches(3.6), Inches(11), Inches(1.0), title, 44, WHITE, True)
        text(s, Inches(1.0), Inches(4.6), Inches(11), Inches(0.8), desc, 22, RGBColor(0xC9, 0xD3, 0xE3))
        return s

    def bullets(self, title, items, notes=None, size=24):
        s = self._new(notes)
        self._chrome(s, title)
        text(s, Inches(0.9), Inches(1.5), Inches(11.6), Inches(5.2), items, size, bullet="•  ", para_space=16)
        return s

    def image_right(self, title, items, image, notes=None, size=24, img_w=5.6):
        """왼쪽 글, 오른쪽 큰 그림"""
        s = self._new(notes)
        self._chrome(s, title)
        iw = Inches(img_w)
        picture_cover(s, self.img(image), SW - iw - Inches(0.5), Inches(1.4), iw, Inches(5.45))
        text(s, Inches(0.9), Inches(1.6), SW - iw - Inches(1.8), Inches(5.2), items, size,
             bullet="•  ", para_space=18)
        self._credit(s, [image])
        return s

    def analogy(self, title, image, big, items, notes=None):
        """비유 슬라이드: 왼쪽 큰 사진, 오른쪽 '비유' 라벨 + 큰 문장 + 설명"""
        s = self._new(notes)
        self._chrome(s, title)
        picture_cover(s, self.img(image), Inches(0.5), Inches(1.4), Inches(6.2), Inches(5.45))
        tag = rect(s, Inches(7.1), Inches(1.55), Inches(1.2), Inches(0.45), ORANGE,
                   shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        text(s, Inches(7.1), Inches(1.55), Inches(1.2), Inches(0.45), "비유", 16, WHITE, True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(7.1), Inches(2.15), Inches(5.8), Inches(1.6), big, 28, INK, True, spacing=1.1)
        text(s, Inches(7.1), Inches(3.9), Inches(5.8), Inches(3.0), items, 21, bullet="•  ", para_space=12)
        self._credit(s, [image])
        return s

    def two_images(self, title, left, right, notes=None):
        """left/right: (image, 라벨, [설명])"""
        s = self._new(notes)
        self._chrome(s, title)
        for i, (image, label, items) in enumerate((left, right)):
            x = Inches(0.5 + i * 6.32)
            picture_cover(s, self.img(image), x, Inches(1.4), Inches(6.0), Inches(3.3))
            col = BLUE if i == 0 else ORANGE
            rect(s, x, Inches(4.8), Inches(6.0), Inches(0.6), col)
            text(s, x, Inches(4.8), Inches(6.0), Inches(0.6), label, 22, WHITE, True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            text(s, x + Inches(0.2), Inches(5.5), Inches(5.6), Inches(1.4), items, 19, bullet="•  ", para_space=6)
        self._credit(s, [left[0], right[0]], x=Inches(2.0), w=Inches(9.9))
        return s

    def video(self, title, video, poster_png, side=None, notes=None):
        """manim 영상 슬라이드. side가 있으면 오른쪽에 핵심 3줄"""
        s = self._new(notes)
        self._chrome(s, title)
        poster(video, poster_png)
        if side:
            w = Inches(9.0)
            h = Emu(int(w * 9 / 16))
            x, y = Inches(0.5), Inches(1.45)
            text(s, Inches(9.8), Inches(1.6), Inches(3.1), Inches(5.0), side, 19, bullet="▸ ", para_space=18)
        else:
            h = Inches(5.45)
            w = Emu(int(h * 16 / 9))
            x, y = Emu(int((SW - w) / 2)), Inches(1.4)
        frame = rect(s, x - Inches(0.04), y - Inches(0.04), w + Inches(0.08), h + Inches(0.08), WHITE, line=LIGHT)
        s.shapes.add_movie(video, x, y, w, h, poster_frame_image=poster_png, mime_type="video/mp4")
        text(s, x, y + h + Inches(0.05), w, Inches(0.3), "▶ 영상을 클릭하면 재생", 11, GRAY)
        return s

    def definition(self, title, term, body, items, image=None, notes=None):
        s = self._new(notes)
        self._chrome(s, title)
        right = Inches(5.0) if image else 0
        bw = SW - Inches(1.0) - right - (Inches(0.3) if image else 0)
        rect(s, Inches(0.5), Inches(1.5), bw, Inches(2.1), WHITE, line=BLUE)
        rect(s, Inches(0.5), Inches(1.5), Inches(0.12), Inches(2.1), BLUE)
        text(s, Inches(0.85), Inches(1.6), bw - Inches(0.5), Inches(0.6), term, 22, BLUE, True)
        text(s, Inches(0.85), Inches(2.2), bw - Inches(0.5), Inches(1.4), body, 26, INK, True)
        text(s, Inches(0.8), Inches(3.95), bw - Inches(0.3), Inches(3.0), items, 21, bullet="•  ", para_space=12)
        if image:
            picture_cover(s, self.img(image), SW - right - Inches(0.5), Inches(1.5), right, Inches(5.35))
            self._credit(s, [image])
        return s

    def table(self, title, header, rows, notes=None, col_w=None, top=1.5, size=18, extra=None):
        s = self._new(notes)
        self._chrome(s, title)
        nr, nc = len(rows) + 1, len(header)
        w = Inches(12.3)
        shp = s.shapes.add_table(nr, nc, Inches(0.5), Inches(top), w, Inches(0.6 * nr))
        tbl = shp.table
        if col_w:
            for i, cw in enumerate(col_w):
                tbl.columns[i].width = Inches(cw)
        for r in range(nr):
            for c in range(nc):
                cell = tbl.cell(r, c)
                val = header[c] if r == 0 else rows[r - 1][c]
                cell.text = ""
                p = cell.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                run = p.add_run()
                run.text = val
                head = r == 0 or (c == 0 and header[0] == "")
                _font(run, size, WHITE if r == 0 else INK, head)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = INK if r == 0 else (LIGHT if c == 0 else WHITE)
        if extra:
            text(s, Inches(0.7), Inches(top + 0.6 * nr + 0.3), Inches(12), Inches(1.2), extra, 20)
        return s

    def cards(self, title, cards, notes=None, footer_text=None):
        """cards: [(제목, [내용], 색)] 2~4개"""
        s = self._new(notes)
        self._chrome(s, title)
        n = len(cards)
        gap = Inches(0.3)
        cw = Emu(int((SW - Inches(1.0) - gap * (n - 1)) / n))
        for i, (head, items, col) in enumerate(cards):
            x = Inches(0.5) + (cw + gap) * i
            rect(s, x, Inches(1.5), cw, Inches(4.9), WHITE, line=LIGHT)
            rect(s, x, Inches(1.5), cw, Inches(0.8), col)
            text(s, x, Inches(1.5), cw, Inches(0.8), head, 22, WHITE, True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            text(s, x + Inches(0.25), Inches(2.5), cw - Inches(0.5), Inches(3.8), items, 19,
                 bullet="•  ", para_space=12)
        if footer_text:
            text(s, Inches(0.5), Inches(6.45), Inches(12.3), Inches(0.5), footer_text, 18, ORANGE, True,
                 align=PP_ALIGN.CENTER)
        return s

    def problem(self, title, tag, question, items=None, image=None, notes=None):
        s = self._new(notes)
        self._chrome(s, title)
        right = Inches(4.6) if image else 0
        w = SW - Inches(1.0) - right - (Inches(0.3) if image else 0)
        rect(s, Inches(0.5), Inches(1.5), Inches(1.6), Inches(0.5), BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        text(s, Inches(0.5), Inches(1.5), Inches(1.6), Inches(0.5), tag, 16, WHITE, True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(0.5), Inches(2.2), w, Inches(1.8), question, 24, INK, True, spacing=1.25)
        if items:
            text(s, Inches(0.7), Inches(4.1), w - Inches(0.2), Inches(2.8), items, 21, para_space=10)
        if image:
            picture_cover(s, self.img(image), SW - right - Inches(0.5), Inches(1.5), right, Inches(5.35))
            self._credit(s, [image])
        return s

    def save(self, path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.prs.save(path)
        return path
